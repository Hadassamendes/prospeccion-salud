#!/usr/bin/env python3
"""Limpia los resultados del scraper y los agrega al Excel de prospectos.

Uso:  python3 limpiar.py datos/resultados_YYYY-MM-DD.csv

- Quita duplicados (contra todo lo que ya está en el Excel) y negocios cerrados.
- Marca: sin página web, solo redes sociales/directorio, calificación baja, pocas reseñas.
- Detecta celulares colombianos y arma el enlace de WhatsApp.
- Calcula una prioridad y conserva las columnas Estado y Notas que tú llenas a mano.
"""
import datetime
import json
import re
import shutil
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = Path(__file__).resolve().parent
EXCEL = BASE / "Prospectos.xlsx"
CONFIG = json.loads((BASE / "config.json").read_text(encoding="utf-8"))

COLUMNAS = [
    "Prioridad", "Puntaje", "Oportunidades", "Negocio", "Nicho", "Ciudad", "Zona",
    "Categoría Maps", "Teléfono", "Tipo teléfono", "WhatsApp", "Web", "Estado web",
    "Calificación", "Reseñas", "Dirección", "Email", "Google Maps", "Encontrado",
    "Estado", "Notas", "ID",
]
ESTADOS = ["Nuevo", "Contactado", "Respondió", "Reunión", "Propuesta", "Cliente", "No interesado", "No contactar"]

REDES = (
    "facebook.", "fb.com", "instagram.", "linktr.ee", "wa.me", "whatsapp.", "api.whatsapp",
    "tiktok.", "twitter.", "x.com", "youtube.", "doctoralia.", "doctoranytime.", "topdoctors.",
    "business.site", "sites.google.", "g.page", "google.com", "linkedin.", "beacons.ai", "bio.link",
)
CERRADO = ("closed", "cerrado")


# ---------------------------------------------------------------- utilidades
def col(df: pd.DataFrame, *nombres: str) -> pd.Series:
    """Devuelve la primera columna que exista (el scraper cambia nombres entre versiones)."""
    for n in nombres:
        if n in df.columns:
            return df[n].fillna("").astype(str).str.strip()
    return pd.Series([""] * len(df), index=df.index)


def a_numero(valor: str) -> float | None:
    valor = str(valor).replace(",", ".").strip()
    try:
        return float(valor)
    except ValueError:
        return None


def clasificar_telefono(tel: str) -> tuple[str, str, str]:
    """Devuelve (teléfono legible, tipo, enlace WhatsApp)."""
    digitos = re.sub(r"\D", "", tel)
    if digitos.startswith("57") and len(digitos) == 12:
        digitos = digitos[2:]
    if len(digitos) == 10 and digitos.startswith("3"):
        legible = f"+57 {digitos[:3]} {digitos[3:6]} {digitos[6:]}"
        return legible, "Celular", f"https://wa.me/57{digitos}"
    if len(digitos) == 10 and digitos.startswith("60"):
        return f"+57 {digitos[:3]} {digitos[3:6]} {digitos[6:]}", "Fijo", ""
    if digitos:
        return tel, "Otro", ""
    return "", "Sin teléfono", ""


def estado_web(web: str) -> str:
    w = web.lower()
    if not w:
        return "Sin web"
    if any(r in w for r in REDES):
        return "Solo redes/directorio"
    return "Tiene web"


def calcular(fila: dict) -> tuple[int, str, str]:
    razones, puntos = [], 0
    if fila["Estado web"] == "Sin web":
        razones.append("Sin página web"); puntos += 3
    elif fila["Estado web"] == "Solo redes/directorio":
        razones.append("Solo redes/directorio"); puntos += 2

    cal, res = fila["Calificación"], fila["Reseñas"]
    if cal is not None and res and cal < CONFIG["calificacion_baja"]:
        razones.append(f"Calificación baja ({cal:.1f})"); puntos += 2
    if res is None or res < CONFIG["pocas_resenas"]:
        razones.append(f"Pocas reseñas ({int(res or 0)})"); puntos += 1

    if fila["Tipo teléfono"] == "Celular":
        puntos += 1
    elif fila["Tipo teléfono"] == "Sin teléfono":
        puntos -= 2

    prioridad = "Alta" if puntos >= 4 else "Media" if puntos >= 2 else "Baja"
    return puntos, prioridad, " · ".join(razones) or "—"


# ---------------------------------------------------------------- limpieza
def limpiar(csv: Path) -> pd.DataFrame:
    crudo = pd.read_csv(csv, dtype=str, keep_default_na=False, on_bad_lines="skip")
    if crudo.empty:
        return pd.DataFrame(columns=COLUMNAS)

    meta = col(crudo, "input_id").str.split("|", expand=True).reindex(columns=[0, 1, 2]).fillna("").apply(lambda s: s.str.strip())
    df = pd.DataFrame({
        "Negocio": col(crudo, "title"),
        "Ciudad": meta[0], "Nicho": meta[1], "Zona": meta[2],
        "Categoría Maps": col(crudo, "category"),
        "TelCrudo": col(crudo, "phone"),
        "Web": col(crudo, "website", "web_site"),
        "CalCruda": col(crudo, "review_rating", "rating"),
        "ResCrudas": col(crudo, "review_count", "reviews"),
        "Dirección": col(crudo, "address", "complete_address"),
        "Email": col(crudo, "emails", "email").str.strip("[]").str.replace('"', ""),
        "Google Maps": col(crudo, "link"),
        "Status": col(crudo, "status"),
        "cid": col(crudo, "cid", "data_id"),
    })

    df = df[df["Negocio"] != ""]
    df = df[~df["Status"].str.lower().str.contains("|".join(CERRADO), na=False)]

    df["ID"] = df["cid"].where(df["cid"] != "", df["Google Maps"])
    sin_id = df["ID"] == ""
    df.loc[sin_id, "ID"] = (df["Negocio"].str.lower() + "|" + df["TelCrudo"].str.replace(r"\D", "", regex=True))[sin_id]
    df = df.drop_duplicates("ID")

    tel = df["TelCrudo"].apply(clasificar_telefono)
    df["Teléfono"] = tel.str[0]
    df["Tipo teléfono"] = tel.str[1]
    df["WhatsApp"] = tel.str[2]
    df["Estado web"] = df["Web"].apply(estado_web)
    df["Calificación"] = df["CalCruda"].apply(a_numero)
    df["Reseñas"] = df["ResCrudas"].apply(a_numero)

    calc = df.apply(lambda f: calcular(f.to_dict()), axis=1)
    df["Puntaje"] = calc.str[0]
    df["Prioridad"] = calc.str[1]
    df["Oportunidades"] = calc.str[2]
    df["Encontrado"] = datetime.date.today().isoformat()
    df["Estado"] = "Nuevo"
    df["Notas"] = ""
    return df[COLUMNAS]


# ---------------------------------------------------------------- Excel
def leer_existente() -> pd.DataFrame:
    if not EXCEL.exists():
        return pd.DataFrame(columns=COLUMNAS)
    # Se lee con openpyxl para recuperar las URLs detrás de "Abrir chat" / "Ver ficha".
    ws = load_workbook(EXCEL)["Prospectos"]
    filas = list(ws.iter_rows())
    if not filas:
        return pd.DataFrame(columns=COLUMNAS)
    encabezados = [str(c.value or "") for c in filas[0]]
    registros = []
    for fila in filas[1:]:
        registro = {}
        for nombre, celda in zip(encabezados, fila):
            valor = celda.hyperlink.target if celda.hyperlink is not None else celda.value
            registro[nombre] = "" if valor is None else str(valor)
        if any(registro.values()):
            registros.append(registro)
    previo = pd.DataFrame(registros, columns=encabezados)
    for c in COLUMNAS:
        if c not in previo.columns:
            previo[c] = ""
    for c in ("Puntaje", "Calificación", "Reseñas"):
        previo[c] = pd.to_numeric(previo[c], errors="coerce")
    return previo[COLUMNAS]


ORDEN_PRIORIDAD = {"Alta": 0, "Media": 1, "Baja": 2}
COLORES = {"Alta": "F8D7DA", "Media": "FFF3CD", "Baja": "E9ECEF"}
ANCHOS = {"Prioridad": 10, "Puntaje": 8, "Oportunidades": 42, "Negocio": 34, "Nicho": 18, "Ciudad": 12,
          "Zona": 14, "Categoría Maps": 22, "Teléfono": 17, "Tipo teléfono": 12, "WhatsApp": 12, "Web": 28,
          "Estado web": 20, "Calificación": 11, "Reseñas": 9, "Dirección": 40, "Email": 26,
          "Google Maps": 12, "Encontrado": 12, "Estado": 14, "Notas": 30, "ID": 10}


def escribir_hoja(ws, df: pd.DataFrame, con_validacion: bool) -> None:
    ws.append(COLUMNAS)
    for fila in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in fila])

    encabezado = PatternFill("solid", fgColor="1F3A5F")
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = encabezado
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions

    for i, nombre in enumerate(COLUMNAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ANCHOS.get(nombre, 14)

    idx = {n: i + 1 for i, n in enumerate(COLUMNAS)}
    for r in range(2, ws.max_row + 1):
        for nombre, texto in (("WhatsApp", "Abrir chat"), ("Google Maps", "Ver ficha"), ("Web", None)):
            celda = ws.cell(row=r, column=idx[nombre])
            url = str(celda.value or "")
            if url.startswith("http"):
                celda.hyperlink = url
                celda.font = Font(color="0563C1", underline="single")
                if texto:
                    celda.value = texto
        ws.cell(row=r, column=idx["Calificación"]).number_format = "0.0"

    ultima = max(ws.max_row, 2)
    letra_p = get_column_letter(idx["Prioridad"])
    for prioridad, color in COLORES.items():
        ws.conditional_formatting.add(
            f"A2:{get_column_letter(len(COLUMNAS))}{ultima}",
            FormulaRule(formula=[f'${letra_p}2="{prioridad}"'], fill=PatternFill("solid", fgColor=color)),
        )
    ws.column_dimensions[get_column_letter(idx["ID"])].hidden = True

    if con_validacion:
        letra_e = get_column_letter(idx["Estado"])
        dv = DataValidation(type="list", formula1='"' + ",".join(ESTADOS) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{letra_e}2:{letra_e}{ultima + 5000}")


def escribir_resumen(ws, todo: pd.DataFrame, nuevos: pd.DataFrame) -> None:
    negrita = Font(bold=True)
    ws.append(["Resumen de prospección"]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Actualizado: {datetime.datetime.now():%Y-%m-%d %H:%M}"])
    ws.append([])
    ws.append(["Total en la base", len(todo)])
    ws.append(["Nuevos en la última corrida", len(nuevos)])
    ws.append(["Nuevos con prioridad Alta", int((nuevos["Prioridad"] == "Alta").sum())])
    ws.append(["Con WhatsApp (celular)", int((todo["Tipo teléfono"] == "Celular").sum())])
    ws.append(["Sin página web", int((todo["Estado web"] == "Sin web").sum())])
    ws.append([])

    for titulo, campo in (("Por ciudad", "Ciudad"), ("Por nicho", "Nicho"), ("Por estado del embudo", "Estado")):
        ws.append([titulo, "Total", "Alta", "Media", "Baja"])
        for c in ws[ws.max_row]:
            c.font = negrita
        tabla = todo.groupby(campo)["Prioridad"].value_counts().unstack(fill_value=0)
        for nombre, fila in tabla.iterrows():
            ws.append([nombre or "(sin dato)", int(fila.sum()), int(fila.get("Alta", 0)),
                       int(fila.get("Media", 0)), int(fila.get("Baja", 0))])
        ws.append([])
    ws.column_dimensions["A"].width = 32
    for letra in "BCDE":
        ws.column_dimensions[letra].width = 10


def guardar(todo: pd.DataFrame, nuevos: pd.DataFrame) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Nuevos hoy"
    escribir_hoja(ws, nuevos, con_validacion=False)
    escribir_hoja(wb.create_sheet("Prospectos"), todo, con_validacion=True)
    escribir_resumen(wb.create_sheet("Resumen"), todo, nuevos)

    if EXCEL.exists():
        respaldos = BASE / "datos" / "respaldos"
        respaldos.mkdir(parents=True, exist_ok=True)
        shutil.copy2(EXCEL, respaldos / f"Prospectos_{datetime.datetime.now():%Y%m%d_%H%M}.xlsx")

    temporal = EXCEL.with_suffix(".tmp.xlsx")
    wb.save(temporal)
    temporal.replace(EXCEL)


def ordenar(df: pd.DataFrame) -> pd.DataFrame:
    return df.assign(_o=df["Prioridad"].map(ORDEN_PRIORIDAD).fillna(3)) \
             .sort_values(["_o", "Puntaje", "Reseñas"], ascending=[True, False, True]) \
             .drop(columns="_o")


def main() -> None:
    if len(sys.argv) < 2:
        sys.exit("Uso: python3 limpiar.py <archivo_resultados.csv>")
    csv = Path(sys.argv[1])
    if not csv.exists() or csv.stat().st_size == 0:
        sys.exit(f"No hay resultados en {csv}")

    hoy = limpiar(csv)
    previo = leer_existente()
    nuevos = ordenar(hoy[~hoy["ID"].isin(set(previo["ID"]))])
    todo = ordenar(pd.concat([previo, nuevos], ignore_index=True))
    guardar(todo, nuevos)

    altas = int((nuevos["Prioridad"] == "Alta").sum())
    print(f"Leídos: {len(hoy)} | Nuevos: {len(nuevos)} (Alta: {altas}) | Total en la base: {len(todo)}")
    (BASE / "datos").mkdir(exist_ok=True)
    (BASE / "datos" / "ultimo_resumen.txt").write_text(f"{len(nuevos)} nuevos, {altas} de prioridad alta", encoding="utf-8")


if __name__ == "__main__":
    main()
